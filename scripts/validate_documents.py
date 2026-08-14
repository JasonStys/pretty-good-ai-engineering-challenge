"""Structural validation for the committed Word and Excel engineering reports."""

from __future__ import annotations

import argparse
import zipfile
from pathlib import Path

from openpyxl import load_workbook

EXPECTED_DOCX = {"engineering-report.docx", "test-validation-report.docx"}
EXPECTED_XLSX = {"engineering-metrics.xlsx", "test-results.xlsx"}
REQUIRED_SHEETS = {
    "engineering-metrics.xlsx": {"Performance", "Complexity", "Resource_Model", "Scenarios"},
    "test-results.xlsx": {"Summary", "Test_Groups", "Security", "CI_Workflows"},
}


def validate_docx(path: Path) -> list[str]:
    """Check the required OPC parts and meaningful document text."""
    errors = []
    try:
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
            required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
            missing = required - names
            if missing:
                errors.append(f"{path}: missing DOCX parts {sorted(missing)}")
            document_xml = package.read("word/document.xml")
            if len(document_xml) < 5_000:
                errors.append(f"{path}: document body is unexpectedly small")
    except (OSError, KeyError, zipfile.BadZipFile) as exc:
        errors.append(f"{path}: invalid DOCX package: {exc}")
    return errors


def validate_xlsx(path: Path) -> list[str]:
    """Check workbook structure, formulas, and cached error cells."""
    errors = []
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        return [f"{path}: invalid workbook: {exc}"]
    missing = REQUIRED_SHEETS[path.name] - set(workbook.sheetnames)
    if missing:
        errors.append(f"{path}: missing worksheets {sorted(missing)}")
    formula_cells: list[tuple[str, str]] = []
    error_count = 0
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.append((worksheet.title, cell.coordinate))
                error_count += cell.data_type == "e"
    if len(formula_cells) < 3:
        errors.append(f"{path}: expected at least three formulas; found {len(formula_cells)}")
    if error_count:
        errors.append(f"{path}: contains {error_count} formula error cells")
    workbook.close()
    calculated = load_workbook(path, data_only=True, read_only=True)
    missing_cache = [
        f"{sheet}!{coordinate}"
        for sheet, coordinate in formula_cells
        if calculated[sheet][coordinate].value is None
    ]
    formula_errors = [
        f"{sheet}!{coordinate}"
        for sheet, coordinate in formula_cells
        if calculated[sheet][coordinate].data_type == "e"
        or str(calculated[sheet][coordinate].value).startswith("#")
    ]
    if missing_cache:
        errors.append(f"{path}: formulas missing cached results: {missing_cache}")
    if formula_errors:
        errors.append(f"{path}: formulas contain calculated errors: {formula_errors}")
    calculated.close()
    return errors


def main() -> int:
    """Validate every required report and return a CI-friendly exit status."""
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path, nargs="?", default=Path("docs"))
    args = parser.parse_args()
    errors = []
    for name in sorted(EXPECTED_DOCX):
        path = args.directory / name
        errors.extend(validate_docx(path) if path.exists() else [f"{path}: missing"])
    for name in sorted(EXPECTED_XLSX):
        path = args.directory / name
        errors.extend(validate_xlsx(path) if path.exists() else [f"{path}: missing"])
    if errors:
        print("Document validation failed:")
        print("\n".join(f"- {error}" for error in errors))
        return 1
    print("Validated 2 DOCX reports and 2 XLSX workbooks; no structural errors found.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
